from pathlib import Path
import json
import csv
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
ASSETS = OUT / 'report-assets'
ASSETS.mkdir(exist_ok=True)
TEMP = Path(r'C:\Users\navee\AppData\Local\Temp')

def crop(name, source, box):
    if (ASSETS / (name + '.png')).exists():
        return ASSETS / (name + '.png')
    im = Image.open(TEMP / source)
    im.crop(box).save(ASSETS / (name + '.png'))
    return ASSETS / (name + '.png')

five = crop('five-users', 'codex-clipboard-9ad5b847-7d5a-448e-afbd-9311e818c085.png', (293, 252, 1147, 337))
ten = crop('ten-users', 'codex-clipboard-87c65d96-b247-4aa1-830b-390b51851e6d.png', (294, 253, 1145, 337))
latest = crop('latest-statistics', 'codex-clipboard-45f438b0-c7c2-4112-9567-a44a3a40f055.png', (177, 536, 1898, 670))
graph = crop('five-second-graph', 'codex-clipboard-4350e244-e386-4a77-979e-0291841112da.png', (374, 355, 1565, 890))
failure = crop('authentication-failure', 'codex-clipboard-02a8975f-ee2b-4c6e-8589-1fef4db1975d.png', (177, 536, 1897, 798))
auth = crop('environment-token-check', 'codex-clipboard-04448421-eac6-45c3-90e8-0408cca646ab.png', (12, 610, 577, 723))
azure = crop('azure-server-errors', 'codex-clipboard-2c616f54-0970-4f4a-a444-8fbaa5ea00ae.png', (362, 337, 1890, 938))

def evidence_crop(name, path, box):
    if (ASSETS / (name + '.png')).exists():
        return ASSETS / (name + '.png')
    Image.open(path).crop(box).save(ASSETS / (name + '.png'))
    return ASSETS / (name + '.png')

create_run = ROOT / 'jmeter-results' / 'create-load-20261018-145949'
notification_run = ROOT / 'jmeter-results' / 'notification-load-20261018-154729'
correlation_run = ROOT / 'jmeter-results' / 'notification-correlation-20261018-152055-028'
write_image = evidence_crop('create-read-results', create_run / 'Create-and-Read-Order-Terminal.png', (0, 95, 1390, 250))
notification_image = evidence_crop('notification-load-results', notification_run / 'Notification-Load-Terminal.png', (0, 383, 870, 458))
correlation_image = evidence_crop('notification-correlation', correlation_run / 'Order-Notification-Correlation-Terminal.png', (0, 59, 900, 300))
inbox_image = evidence_crop('inbox-batch-orders', correlation_run / 'Inbox-Order-Notifications.jpeg', (225, 84, 1590, 463))
email_image = evidence_crop('opened-order-email', correlation_run / 'Opened-Order-Notification.jpeg', (686, 210, 1226, 765))

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.65)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.05
styles['Title'].font.size = Pt(23)
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].paragraph_format.space_before = Pt(8)
styles['Heading 2'].font.size = Pt(12)
styles['Caption'].font.size = Pt(9)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Sellora Sprint 3 load testing  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text): return doc.add_paragraph(text)
def h(text): return doc.add_heading(text, 1)
def pic(path, caption, width=7.15):
    para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(3)
    para.add_run().add_picture(str(path), width=Inches(width))
    doc.add_paragraph(caption, 'Caption')
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for c, text in zip(t.rows[0].cells, headers): c.text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j]); cell.vertical_alignment = 1
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            pr.append(borders)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE6F1' if i == 0 else ('F5F7FA' if i % 2 == 0 else 'FFFFFF')); pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                m=OxmlElement('w:'+side); m.set(qn('w:w'),'70'); m.set(qn('w:type'),'dxa'); margins.append(m)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                if j: para.alignment=1
                for r in para.runs: r.font.size=Pt(10); r.bold=(i==0)
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_paragraph('Sellora Sprint 3 Load Testing Report', 'Title')
p('SE3112 Advanced Software Engineering | Load and Performance Testing')
p('Test dates 17 and 18 October 2026 | Report date 18 October 2026')
h('Purpose and outcome')
p('I tested Sprint 3 Order and Notification behaviour using Apache JMeter 5.6.3. Order List passed at up to 50 direct Azure clients in one measured run, with inconsistent repeats. A ten-order write batch succeeded but missed the proposed p95 target. Notification List passed at five clients; all ten orders had notification records with both recipients marked Sent. Reliable maximum capacity was not established.')
h('Test scope and setup')
p('Gateway tests passed through WSO2 API Manager; direct Azure tests bypassed it. Order tests used a Sales Rep token and Notification tests used a Company Admin token. Clients shared one account per role; this did not represent distinct user identities.')
p('Read paths were /api/orders and /api/notifications, with pagination. Creation used POST /api/orders followed by GET /api/orders/{orderId}. Read assertions checked status and response structure; the created-order GET also checked the extracted order ID. Detailed workload differences appear with each result.')
table(['Setting', 'Configuration'], [
('Workload','Up to 50 concurrent clients tested'),
('Exploratory runs','60 seconds including a 5 second ramp up'),
('Capacity runs','180 seconds with 10 second ramp up; first 30 seconds excluded'),
('Additional service coverage','10 orders created and retrieved; Notification read load at 5 clients'),
('Pause between read iterations','1,000 ms per client; no pause added to the finite write batch'),
('Execution','GUI exploration followed by command line runs'),
('Proposed acceptance targets','p95 at most 2,000 ms and errors below 1%'),
], [2.15,5.05])
p('The targets are provisional assessment targets, not an agreed service SLA. Measured times include the client network, gateway and backend path. These tests do not isolate backend processing time.')
h('Contribution completed')
p('Saved plans use environment-supplied role tokens. Scripts preserve results, generate reports, evaluate targets and correlate order events with notification records. Azure metrics corroborated one error burst; inbox screenshots corroborated email arrival. Integration with the wider group pipeline remains outstanding.')

doc.add_page_break()
h('Exploratory results in JMeter')
p('The following GUI runs established a working authenticated test and introduced concurrency and percentile measurements. These are short exploratory observations, not evidence of the system capacity limit.')
table(['Metric','5 clients','10 clients'], [
('Requests',154,359),('Average ms',891,582),('Median ms',651,482),('p95 ms',2065,1103),('p99 ms',2920,2138),('Maximum ms',4301,4783),('Throughput requests per second','2.5','6.1'),('Error rate','0%','0%')
], [3.3,1.95,1.95])
pic(five, 'Figure 1  Five client Aggregate Report from the 60 second GUI run')
pic(ten, 'Figure 2  Ten client Aggregate Report from the 60 second GUI run')
h('Interpretation')
p('Throughput increased in the ten-client observation, with no request or assertion failures. The lower average at the higher client count does not prove that added load improved performance; network conditions, connection reuse and server state can vary between runs.')
p('A successful response can still be slow. Percentiles describe the slower part of the response distribution and complement the average. A p95 of 1,103 ms means approximately 95% of measured responses completed at or below that time.')

doc.add_page_break()
h('Command line results and performance targets')
p('Two successful authenticated command line runs used the intended ten-client workload. Both included startup. The later run was slower despite the same intended workload, so the evidence does not establish consistent performance.')
rows=[]
for label, folder in [('Earlier run','20261018-111639'),('Latest run','orders-10users-20261018-114507')]:
    s=json.loads((ROOT/'jmeter-results'/folder/'html'/'statistics.json').read_text())['Total']
    rows.append((label,str(s['sampleCount']),f"{s['meanResTime']:.1f}",f"{s['pct2ResTime']:.1f}",f"{s['throughput']:.2f}",'0%','PASS' if s['pct2ResTime']<=2000 else 'FAIL'))
table(['Run','Requests','Avg ms','p95 ms','Req per sec','Errors','Gate'], rows,[1.25,.7,.95,.95,.9,.65,.8])
pic(latest, 'Figure 3  Latest HTML dashboard statistics showing 257 requests and zero failures')
p('The latest run had p99 13,042.46 ms and maximum 15,357 ms. It missed the proposed p95 target while meeting the error target. The dashboard request pass percentage checks request success; it is separate from the performance gate.')
pic(graph, 'Figure 4  Earlier run response time averages grouped into five second intervals')
p('Figure 4 belongs to the earlier 343-request run. It shows a higher initial interval and temporary later peaks, followed by lower averages. Each point averages multiple requests. It does not reveal the cause of delay or prove a sustained upward trend.')

doc.add_page_break()
h('Authentication issue and resolution')
p('An initial command line run returned 401 Unauthorized for all 432 requests. Those timings represent rejected requests and were excluded from successful-operation comparisons. A fresh token and saved plan were used before rerunning successfully; the precise reason the previous token was rejected was not independently established.')
pic(failure, 'Figure 5  Failed run showing 100 percent errors and 401 Unauthorized responses')
p('The token was subsequently moved outside the test plan. JMeter reads SELLORA_ACCESS_TOKEN from the launching PowerShell process. A short verification request returned HTTP 200 with success true.')
pic(auth, 'Figure 6  Successful verification after externalizing the access token', width=5.4)
h('Automation and saved evidence')
p('List Orders.jmx accepts users, rampup, duration and thinkTime properties. Check-Performance.ps1 reads the HTML report statistics using the default 90, 95 and 99 percentile settings. Its exit codes are 0 for targets met, 1 for targets missed and 2 for invalid input. The checker was verified against both successful-operation reports: the earlier report passed and the latest failed.')
p('Evidence folders: jmeter-results/20261018-111639 and jmeter-results/orders-10users-20261018-114507. The authentication failure is retained in jmeter-results/20261018-110949. Preserve full report folders together with requests.jtl and jmeter.log.')
h('Limitations and remaining work')
p('Coverage includes Order reads, a small scheduled-order create-and-retrieve batch, Notification reads and stored notification completion records. Payment, login capacity and maximum notification-processing capacity were not tested. Clients shared one account per role. Azure error counts corroborated a failure but did not establish its cause. No whole-system or production maximum is claimed.')
p('The plans have a 5,000 ms connect timeout and a 10,000 ms response timeout. These were added after the early exploratory runs and were present for the capacity runs. They are not strict total transaction deadlines. The wider group pipeline still needs to invoke the runner and preserve its exit status and evidence.')
p('Screenshots in this report were supplied during testing and cropped only for readability. Numeric command line results were checked against the saved JMeter statistics files.')

doc.add_page_break()
h('Capacity exploration results')
p('Each capacity run lasted 180 seconds, with a 10 second ramp up and a 1 second pause per client. Requests starting in the first 30 seconds were excluded from the measurement report. The tables below use only the measured window; full raw results retain warmup requests and failures.')
capacity_folders = [
('1','capacity-direct-azure-1users-20261018-123901-028'),
('5','capacity-direct-azure-5users-20261018-133628-198'),
('10','capacity-direct-azure-10users-20261018-134021-776'),
('20','capacity-direct-azure-20users-20261018-134408-840'),
('30','capacity-direct-azure-30users-20261018-135151-233'),
('50 first','capacity-direct-azure-50users-20261018-135614-434'),
('50 repeat','capacity-direct-azure-50users-20261018-140636-111')]
capacity_rows=[]
for label,folder in capacity_folders:
    s=json.loads((ROOT/'jmeter-results'/folder/'html'/'statistics.json').read_text())['Total']
    passed=s['pct2ResTime']<=2000 and s['errorPct']<1
    capacity_rows.append((label,str(s['sampleCount']),f"{s['pct2ResTime']:.1f}",f"{s['errorPct']:.2f}%",f"{s['throughput']:.2f}",'PASS' if passed else 'FAIL'))
doc.add_heading('Direct Azure endpoint', 2)
table(['Clients','Requests','p95 ms','Errors','Req per sec','Gate'],capacity_rows,[1.15,1.15,1.25,1.1,1.3,1.05])
p('The direct endpoint is /api/orders on the Azure staging host. Passing at 50 clients demonstrates that the proposed targets were achievable in that measured run; it does not establish consistent capacity or a maximum. The repeat processed 5,505 measured requests at 36.79 requests per second.')
doc.add_heading('Gateway path', 2)
gateway_rows=[]
for label,folder in [('5 first','capacity-5users-20261018-121644-072'),('5 repeat','capacity-5users-20261018-122414-686'),('1','capacity-1users-20261018-123043-280')]:
    s=json.loads((ROOT/'jmeter-results'/folder/'html'/'statistics.json').read_text())['Total']
    gateway_rows.append((label,str(s['sampleCount']),f"{s['pct2ResTime']:.1f}",f"{s['errorPct']:.2f}%",'FAIL'))
table(['Clients','Requests','p95 ms','Errors','Gate'],gateway_rows,[1.3,1.4,1.7,1.4,1.2])
p('Gateway runs showed socket errors or read timeouts even at low concurrency. Direct runs were generally faster, but tests occurred at different times. This comparison narrows investigation without proving that API Manager caused the delays. Direct Azure results cannot substitute for user-facing gateway capacity.')

doc.add_page_break()
h('Azure evidence and final conclusion')
p('The first 50-client direct run took place on 18 October 2026 at approximately 13:56 to 13:59 Sri Lanka time. JMeter recorded 39 HTTP 500 and 32 HTTP 503 responses, for 71 errors. Azure Metrics displayed 71 HTTP server errors in the corresponding test window.')
pic(azure, 'Figure 7  Azure staging metrics showing 71 HTTP server errors in the test window')
p('The matching counts corroborate the observed server failures. CPU Time is measured in seconds of CPU work, not CPU utilization percentage. The combined chart uses different units; the tooltip values are the relevant evidence. No CPU saturation or specific root cause was established.')
h('Why the repeat matters')
p('The repeated 50-client run met both targets in the measured window: p95 870.8 ms, zero errors and 5,505 requests. Its full raw run still contained four HTTP 503 responses during warmup. Reporting measured-window success must not hide those startup failures. The first failed run remains part of the evidence.')
h('Capacity statement')
p('The direct Order List endpoint met the proposed performance targets at up to 50 concurrent clients in one measured run, with all clients sharing a Sales Rep account and pausing one second between iterations. Results at 50 clients were inconsistent across repeats. Reliable maximum capacity was not established, and no capacity claim is made for the whole Sellora system or its gateway path.')
h('Bounded capacity exploration')
p('Capacity exploration stopped after the repeated 50-client workload. No 75-client or higher run was performed. The following sections add a fixed ten-order write batch and one Notification read-load run. These are separate scenarios, not an extension of the capacity search. Further work should investigate failures before making an operational capacity commitment.')

doc.add_page_break()
h('Scheduled order creation and retrieval')
p('A separate finite-loop JMeter plan used five configured Sales Rep clients with a five second ramp up and two iterations each. Each iteration submitted one ScheduledDelivery order with one unit of Maliban Chocolate Biscuit, then retrieved its own created order. No payment or approval request was executed.')
p('The direct Azure staging run completed on 18 October 2026 at approximately 14:59:50 to 14:59:57 Sri Lanka time. All ten POSTs returned 201 and all ten GETs returned 200. The saved ledger contains ten distinct order IDs with reservation IDs, references and returned totals.')
write_stats=json.loads((create_run/'html'/'statistics.json').read_text())
write_rows=[]
for label,key in [('Create Order','POST - Create Order'),('Read Created Order','GET - Read Created Order')]:
    s=write_stats[key]
    write_rows.append((label,str(s['sampleCount']),f"{s['meanResTime']:.1f}",f"{s['pct2ResTime']:.1f}",f"{s['maxResTime']:.0f}",'0%'))
table(['Operation','Requests','Average ms','p95 ms','Max ms','Errors'],write_rows,[2.05,.85,1.15,1.1,1.05,.8])
pic(write_image,'Figure 8  JMeter completed the create and retrieve batch with 20 samples and zero errors')
h('Checks and interpretation')
p('The POST asserted HTTP 201 and the existence of orderId. A JSON Extractor saved the created ID, order reference, reservation ID and total. A conditional controller skipped retrieval after a failed creation. GET asserted HTTP 200 and that its orderId matched the extracted value. Thread-local variables kept each client correlated with its own order.')
p('Creation p95 was 2,422 ms, missing the proposed 2,000 ms target despite zero request or assertion failures. Read p95 was 323 ms and met the target. POST and GET must be evaluated separately: pooling the faster reads with writes would obscure creation latency. With only ten samples per operation, percentile values are descriptive, not a write-capacity estimate.')
h('Data changes and setup issues')
p('These were real staging orders, with database writes, stock reservation or commitment, and OrderPlaced events. The earlier UI example and single JMeter baseline are separate from the ten-order batch. Test records have not been reported as cancelled or cleaned up. Use supported business operations and verify stock reconciliation; do not delete database rows independently.')
p('During setup, an empty POST body caused HTTP 400. A GET assertion initially compared the JSON body with the text 200; correcting it to compare the response code resolved that test configuration error. The existing order was then checked without repeating creation. These setup failures were not counted as failures in the successful concurrent batch.')

doc.add_page_break()
h('Notification API and event processing')
p('Notification reads used a Company Admin token against the direct Azure staging service. A single baseline GET returned 200 in 3,258 ms. The subsequent load run used five clients, a five second ramp up, a 60 second duration and a 1,000 ms pause. The whole run, including startup, was evaluated; it has no 30 second exclusion.')
ns=json.loads((notification_run/'html'/'statistics.json').read_text())['Total']
table(['Requests','Average ms','p95 ms','Maximum ms','Errors','Gate'],[[str(ns['sampleCount']),f"{ns['meanResTime']:.1f}",f"{ns['pct2ResTime']:.1f}",f"{ns['maxResTime']:.0f}",'0%','PASS']],[1.0,1.25,1.25,1.4,1.0,1.2])
pic(notification_image,'Figure 9  Notification List passed the provisional p95 and error targets')
p('The run achieved approximately 3.18 requests per second. The maximum 3,414 ms response does not contradict a p95 of 1,647.4 ms: a small proportion can take longer. This result measures the read API, not the maximum rate of background email dispatch.')
correlations=list(csv.DictReader((correlation_run/'Order-Notification-Correlation.csv').open(encoding='utf-8-sig',newline='')))
completion_times=[float(r['EventToCompletionMs']) for r in correlations]
h('Correlating the existing ten orders')
p('A read-only script looked up notification requests by the ten existing order IDs and selected OrderPlaced. All ten had exactly one matching request, each with Shop and Agency recipients marked Sent: 20 sent-recipient records. The script saved IDs and stored event, intake and dispatch timestamps without generating another order or requesting a resend.')
pic(correlation_image,'Figure 10  Ten matching OrderPlaced requests with both recipients marked Sent')
p(f"Stored event-to-completion times ranged from {min(completion_times)/1000:.2f} to {max(completion_times)/1000:.2f} seconds, averaging {sum(completion_times)/len(completion_times)/1000:.2f} seconds. These values describe this batch, include asynchronous stages, and have no agreed delivery SLA. The snapshot was taken later; it is not a live backlog-drain measurement and does not establish maximum consumer capacity.")

doc.add_page_break()
h('Inbox evidence and final scope')
p('The supplied inbox screenshot shows conversation entries for all ten batch order references. Earlier UI and baseline orders are visible below the batch and were not counted as its results. The opened message for ORD-261018-EGXXMP shows Scheduled delivery, the expected product, quantity one and total LKR 150.00.')
pic(inbox_image,'Figure 11  Inbox entries matching the ten batch order references')
pic(email_image,'Figure 12  Opened email for one order in the JMeter batch',width=3.45)
p('These screenshots corroborate inbox arrival for displayed messages. They do not independently verify both intended recipient addresses, exact inbox-delivery latency or the template claim that messages were sent simultaneously. Recipient outcomes are supported separately by the stored service records.')
p('Completed scope: Order List exploration up to 50 clients, ten concurrent-workflow order creations and retrievals, Notification List at five clients, stored event-processing correlation and inbox evidence. Maximum reliable capacity remains unestablished. Cleanup remains to be verified; group source, slides and pipeline integration are separate submission responsibilities.')

target=OUT/'Sellora_Sprint_3_Load_Testing_Report.docx'
for element in list(doc.styles.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
for element in list(doc.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
doc.save(target)
print(target)
